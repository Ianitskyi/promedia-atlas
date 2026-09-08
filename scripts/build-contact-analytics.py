import json,re
from collections import defaultdict
from pathlib import Path
import openpyxl

source=Path('data/registry-source-2026-08-01.xlsx')
book=openpyxl.load_workbook(source,read_only=True,data_only=True)
main,foreign=(book[s] for s in book.sheetnames)
emails=defaultdict(set); phones=defaultdict(set); addresses=defaultdict(set)

def split_emails(value):
    return re.findall(r'[\w.+-]+@[\w.-]+\.[A-Za-zА-Яа-яІіЇїЄє]{2,}',str(value or ''))
def email_key(v): return v.strip().lower()
def phone_key(v): return re.sub(r'\D','',str(v or ''))[-10:]
def address_key(v): return re.sub(r'\s+',' ',str(v or '')).strip().casefold()
def mask_email(v):
    local,domain=v.split('@',1); return (local[:1]+'***@'+domain) if local else '***@'+domain

for row in main.iter_rows(min_row=3,values_only=True):
    media_id=str(row[5] or '').strip()
    if not media_id: continue
    for email in split_emails(row[14]): emails[email_key(email)].add(media_id)
for row in foreign.iter_rows(min_row=3,values_only=True):
    media_id=str(row[7] or '').strip()
    if not media_id: continue
    for email in split_emails(row[2]): emails[email_key(email)].add(media_id)
    phone=phone_key(row[3]); address=address_key(row[6])
    if len(phone)>=9: phones[phone].add(media_id)
    if len(address)>=12: addresses[address].add(media_id)

def groups(values,kind):
    rows=[]
    for value,ids in values.items():
        if len(ids)<2: continue
        label=mask_email(value) if kind=='email' else ('••• '+value[-4:] if kind=='phone' else 'Спільна адреса')
        rows.append({'label':label,'count':len(ids),'mediaIds':sorted(ids)})
    return sorted(rows,key=lambda x:(-x['count'],x['label']))

email_groups=groups(emails,'email'); phone_groups=groups(phones,'phone'); address_groups=groups(addresses,'address')
out={'date':'01.08.2026','registrationDateAvailable':False,'notes':{
'registrationDate':'У файлі немає окремого поля дати реєстрації, тому часові ряди не розраховуються.',
'addresses':'Повні адреси є лише на аркуші іноземних лінійних медіа. Для українських записів наведені населений пункт та область, але не вулиця й будинок.',
'privacy':'Email і телефони маскуються. Збіг контакту є індикатором для перевірки, а не доказом спільної власності.'},
'emails':{'duplicateGroups':len(email_groups),'mediaInGroups':sum(x['count'] for x in email_groups),'groups':email_groups},
'phones':{'duplicateGroups':len(phone_groups),'mediaInGroups':sum(x['count'] for x in phone_groups),'groups':phone_groups},
'addresses':{'duplicateGroups':len(address_groups),'mediaInGroups':sum(x['count'] for x in address_groups),'groups':address_groups}}
Path('public/data/contact-analytics.json').write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print(json.dumps({k:{'groups':v['duplicateGroups'],'media':v['mediaInGroups']} for k,v in out.items() if k in ('emails','phones','addresses')},ensure_ascii=False))
